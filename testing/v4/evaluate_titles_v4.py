#!/usr/bin/env python3
"""
Technical drawing title + location evaluation pipeline V4.

PURPOSE
-------
V4 keeps all useful V3 analysis and adds:
- practical ("good enough for retrieval") title matching
- practical location matching
- location reference-quality handling
- combined title/location retrieval status
- optional original-file metadata extraction
- grouped summary statistics

INPUT 1: Ground truth (.csv or .xlsx)
Required columns:
    filename
    actual_title
    actual_location
    location_status

The workbook may also contain old predicted_title / predicted_location columns.
V4 deliberately ignores them. Predictions always come from INPUT 2.

INPUT 2: Prediction CSV
Required columns:
    Filename
    Title

Optional location column:
    Location

Optional process metadata:
    Title_confidence
    Title_probability
    Original_size_px
    Original_size_MP
    Crop_size_px
    Sent_size_px
    Scale_factor
    Scale_percent
    Sent_KB
    Preparation_time_s
    First_response_time_s
    Model_time_s
    Total_time_s

INPUT 3: Optional directory containing original .tif/.tiff/.pdf files

DEFAULT PRACTICAL-MATCH RULES
-----------------------------
Title is practically usable when ANY applies:
    - exact match
    - character similarity >= 0.92
    - word F1 similarity >= 0.90

Location is practically usable when ANY applies:
    - exact match
    - normalized exact match
    - character similarity >= 0.85
    - word F1 similarity >= 0.80
    - normalized "core location" tokens contain each other

Location normalization is intentionally more tolerant than title matching:
    - case/punctuation differences are ignored
    - ß is treated as ss
    - common generic stop/station words such as Bahnhof, Bhf., Haltestelle,
      U-Bhf. are ignored for core-token containment

IMPORTANT
---------
"practical match" means usable for retrieval/identification according to the
chosen thresholds. It is NOT the same thing as exact correctness.

Rows whose location_status is "uncertain_by_human" are still shown, but are
excluded from the main reliable location success rate.

Dependencies:
    Standard library
    openpyxl  -> only needed when ground truth is .xlsx
    pillow    -> TIFF metadata (optional)
    pymupdf   -> PDF metadata (optional)

Install optional libraries:
    pip install openpyxl pillow pymupdf
"""

import argparse
import csv
import re
import unicodedata
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path


# ============================================================
# Generic helpers
# ============================================================

def detect_delimiter(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t").delimiter
    except csv.Error:
        return ";"


def read_csv(path):
    delimiter = detect_delimiter(path)
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        if reader.fieldnames is None:
            raise ValueError(f"{path} has no header row.")
        return list(reader)


def read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError(
            "Reading .xlsx ground truth requires openpyxl. "
            "Install it with: pip install openpyxl"
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    iterator = ws.iter_rows(values_only=True)
    headers = next(iterator, None)

    if headers is None:
        raise ValueError(f"{path} is empty.")

    headers = ["" if h is None else str(h).strip() for h in headers]
    rows = []

    for values in iterator:
        row = {}
        for header, value in zip(headers, values):
            if not header:
                continue
            row[header] = "" if value is None else str(value)
        rows.append(row)

    wb.close()
    return rows


def read_groundtruth(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return read_xlsx(path)
    if suffix == ".csv":
        return read_csv(path)
    raise ValueError("Ground truth must be .csv or .xlsx")


def require_columns(rows, required, label):
    if not rows:
        raise ValueError(f"{label} contains no data rows.")

    columns = set(rows[0].keys())
    missing = [c for c in required if c not in columns]

    if missing:
        raise ValueError(
            f"{label} is missing required column(s): {', '.join(missing)}"
        )


def normalize_filename(value):
    if value is None:
        return ""
    return unicodedata.normalize("NFC", str(value)).strip().casefold()


def parse_decimal(value):
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None

    text = text.replace("%", "").replace(",", ".")

    try:
        return float(text)
    except ValueError:
        return None


def average(values):
    values = [v for v in values if v is not None]
    return sum(values) / len(values) if values else None


def safe_round(value, digits=4):
    return "" if value is None else round(value, digits)


def yes_no(value):
    return "YES" if value else "NO"


# ============================================================
# Text metrics
# ============================================================

def character_similarity(actual, predicted):
    return SequenceMatcher(None, actual, predicted).ratio()


def tokenize(text):
    if text is None:
        return []
    return re.findall(r"[0-9A-Za-zÄÖÜäöüß]+", str(text).casefold())


def word_metrics(actual, predicted):
    actual_words = tokenize(actual)
    predicted_words = tokenize(predicted)

    actual_counts = defaultdict(int)
    predicted_counts = defaultdict(int)

    for word in actual_words:
        actual_counts[word] += 1

    for word in predicted_words:
        predicted_counts[word] += 1

    common = sum(
        min(count, predicted_counts.get(word, 0))
        for word, count in actual_counts.items()
    )

    a_total = len(actual_words)
    p_total = len(predicted_words)

    if a_total == 0 and p_total == 0:
        precision = recall = f1 = 1.0
    elif a_total == 0 or p_total == 0:
        precision = recall = f1 = 0.0
    else:
        precision = common / p_total
        recall = common / a_total
        f1 = (
            0.0
            if precision + recall == 0
            else 2 * precision * recall / (precision + recall)
        )

    missing_words = []
    for word, count in actual_counts.items():
        missing_words.extend(
            [word] * max(0, count - predicted_counts.get(word, 0))
        )

    extra_words = []
    for word, count in predicted_counts.items():
        extra_words.extend(
            [word] * max(0, count - actual_counts.get(word, 0))
        )

    return f1, missing_words, extra_words


def normalize_general_text(text):
    """
    Loose normalization used for practical location comparison, not exact title.
    """
    text = unicodedata.normalize("NFKC", str(text or "")).casefold()
    text = text.replace("ß", "ss")
    text = text.replace("straße", "strasse")
    text = re.sub(r"[-_/.,;:()\[\]\"']", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


LOCATION_GENERIC_TOKENS = {
    "u", "u-bhf", "ubhf", "bhf", "bahnhof", "haltestelle",
    "station", "u-bahnhof", "ubahnhof"
}


def core_location_tokens(text):
    normalized = normalize_general_text(text)
    tokens = re.findall(r"[0-9a-zäöü]+", normalized)

    result = []
    for token in tokens:
        token = token.replace("ß", "ss")
        if token not in LOCATION_GENERIC_TOKENS:
            result.append(token)

    return result


def location_core_containment(actual, predicted):
    a = core_location_tokens(actual)
    p = core_location_tokens(predicted)

    if not a or not p:
        return False

    a_set = set(a)
    p_set = set(p)

    # Require at least one meaningful token and full containment in one direction.
    return a_set.issubset(p_set) or p_set.issubset(a_set)


def practical_title_match(
    exact_match,
    char_similarity,
    word_similarity,
    char_threshold,
    word_threshold,
):
    return (
        exact_match
        or char_similarity >= char_threshold
        or word_similarity >= word_threshold
    )


def practical_location_match(
    actual,
    predicted,
    exact_match,
    char_similarity,
    word_similarity,
    char_threshold,
    word_threshold,
):
    if not predicted.strip():
        return False

    normalized_exact = (
        normalize_general_text(actual) == normalize_general_text(predicted)
    )

    return (
        exact_match
        or normalized_exact
        or char_similarity >= char_threshold
        or word_similarity >= word_threshold
        or location_core_containment(actual, predicted)
    )


def suggested_title_error_category(
    actual,
    predicted,
    exact_match,
    char_similarity,
    word_similarity,
    missing_words,
    extra_words,
):
    if exact_match:
        return "CORRECT"

    if not predicted.strip():
        return "NO_TITLE_FOUND"

    a_words = tokenize(actual)
    p_words = tokenize(predicted)

    if char_similarity >= 0.92:
        return "MINOR_OCR_OR_FORMATTING_ERROR"

    if (
        word_similarity >= 0.65
        and len(p_words) < len(a_words)
        and len(missing_words) > len(extra_words)
    ):
        return "PARTIAL_TITLE"

    if (
        word_similarity >= 0.65
        and len(p_words) > len(a_words)
        and len(extra_words) > len(missing_words)
    ):
        return "EXTRA_TEXT"

    if char_similarity >= 0.55 or word_similarity >= 0.50:
        return "MAJOR_RECOGNITION_ERROR"

    return "POSSIBLE_WRONG_REGION_OR_SEVERE_ERROR"


def title_features(text):
    return {
        "title_digit_count": sum(ch.isdigit() for ch in text),
        "title_punctuation_count": sum(
            1 for ch in text if not ch.isalnum() and not ch.isspace()
        ),
        "title_has_umlaut_or_sz": any(ch in "äöüÄÖÜß" for ch in text),
    }


# ============================================================
# Location-reference quality
# ============================================================

def location_reference_quality(status):
    value = str(status or "").strip().casefold()

    if value == "manually_identified":
        return "STRONG"
    if value == "derived_from_title":
        return "DERIVED"
    if value == "uncertain_by_human":
        return "UNCERTAIN"
    return "UNKNOWN"


def location_reference_reliable(status):
    # Derived references are usable for aggregate analysis, but distinguish them
    # from manually identified references in the output.
    return location_reference_quality(status) in {"STRONG", "DERIVED"}


# ============================================================
# Original file metadata
# ============================================================

def blank_file_metadata():
    return {
        "source_file_found": False,
        "file_extension": "",
        "file_size_mb": "",
        "image_width_px": "",
        "image_height_px": "",
        "image_megapixels": "",
        "aspect_ratio": "",
        "dpi_x": "",
        "dpi_y": "",
        "image_mode": "",
        "image_frames": "",
        "pdf_page_count": "",
        "pdf_page_width_pt": "",
        "pdf_page_height_pt": "",
        "pdf_page_width_mm": "",
        "pdf_page_height_mm": "",
        "pdf_embedded_image_count": "",
        "pdf_largest_image_width_px": "",
        "pdf_largest_image_height_px": "",
        "metadata_error": "",
    }


def extract_tiff_metadata(path):
    meta = blank_file_metadata()
    meta["source_file_found"] = True
    meta["file_extension"] = path.suffix.lower()
    meta["file_size_mb"] = round(path.stat().st_size / (1024 * 1024), 4)

    try:
        from PIL import Image

        # These are trusted local university drawings and some exceed Pillow's
        # default pixel safety threshold. We need header metadata from them.
        Image.MAX_IMAGE_PIXELS = None
    except ImportError:
        meta["metadata_error"] = "Pillow not installed"
        return meta

    try:
        with Image.open(path) as img:
            width, height = img.size

            meta["image_width_px"] = width
            meta["image_height_px"] = height
            meta["image_megapixels"] = round(width * height / 1_000_000, 4)
            meta["aspect_ratio"] = round(width / height, 4) if height else ""
            meta["image_mode"] = img.mode
            meta["image_frames"] = getattr(img, "n_frames", 1)

            dpi = img.info.get("dpi")
            if dpi and len(dpi) >= 2:
                meta["dpi_x"] = round(float(dpi[0]), 2)
                meta["dpi_y"] = round(float(dpi[1]), 2)

    except Exception as exc:
        meta["metadata_error"] = f"{type(exc).__name__}: {exc}"

    return meta


def extract_pdf_metadata(path):
    meta = blank_file_metadata()
    meta["source_file_found"] = True
    meta["file_extension"] = path.suffix.lower()
    meta["file_size_mb"] = round(path.stat().st_size / (1024 * 1024), 4)

    try:
        import pymupdf
    except ImportError:
        meta["metadata_error"] = "PyMuPDF not installed"
        return meta

    try:
        doc = pymupdf.open(path)
        meta["pdf_page_count"] = doc.page_count

        total_images = 0
        largest_area = 0
        largest_w = ""
        largest_h = ""

        if doc.page_count > 0:
            rect = doc[0].rect
            meta["pdf_page_width_pt"] = round(rect.width, 2)
            meta["pdf_page_height_pt"] = round(rect.height, 2)
            meta["pdf_page_width_mm"] = round(rect.width * 25.4 / 72, 2)
            meta["pdf_page_height_mm"] = round(rect.height * 25.4 / 72, 2)

        for page_index in range(doc.page_count):
            page = doc[page_index]
            images = page.get_images(full=True)
            total_images += len(images)

            for image in images:
                xref = image[0]
                try:
                    info = doc.extract_image(xref)
                    w = info.get("width")
                    h = info.get("height")

                    if w and h and w * h > largest_area:
                        largest_area = w * h
                        largest_w = w
                        largest_h = h
                except Exception:
                    pass

        meta["pdf_embedded_image_count"] = total_images
        meta["pdf_largest_image_width_px"] = largest_w
        meta["pdf_largest_image_height_px"] = largest_h

        doc.close()

    except Exception as exc:
        meta["metadata_error"] = f"{type(exc).__name__}: {exc}"

    return meta


def build_file_lookup(files_dir):
    lookup = {}

    if files_dir is None:
        return lookup

    files_dir = Path(files_dir)

    if not files_dir.exists():
        raise ValueError(f"Files directory does not exist: {files_dir}")

    for path in files_dir.rglob("*"):
        if not path.is_file():
            continue

        if path.suffix.lower() not in {".tif", ".tiff", ".pdf"}:
            continue

        key = normalize_filename(path.name)

        if key not in lookup:
            lookup[key] = path

    return lookup


def get_file_metadata(filename, lookup):
    if not lookup:
        return blank_file_metadata()

    path = lookup.get(normalize_filename(filename))

    if path is None:
        return blank_file_metadata()

    if path.suffix.lower() in {".tif", ".tiff"}:
        return extract_tiff_metadata(path)

    if path.suffix.lower() == ".pdf":
        return extract_pdf_metadata(path)

    return blank_file_metadata()


# ============================================================
# Bands / grouped analysis
# ============================================================

def title_length_band(length):
    if length <= 30:
        return "<=30"
    if length <= 50:
        return "31-50"
    if length <= 70:
        return "51-70"
    return ">70"


def scale_band(value):
    if value is None:
        return "unknown"
    if value < 25:
        return "<25%"
    if value < 40:
        return "25-40%"
    if value < 60:
        return "40-60%"
    return ">=60%"


def dpi_band(value):
    if value is None:
        return "unknown"
    if value < 150:
        return "<150"
    if value < 300:
        return "150-299"
    if value < 600:
        return "300-599"
    return ">=600"


def megapixel_band(value):
    if value is None:
        return "unknown"
    if value < 10:
        return "<10MP"
    if value < 30:
        return "10-30MP"
    if value < 100:
        return "30-100MP"
    return ">=100MP"


def summarize_groups(results, field):
    grouped = defaultdict(list)

    for row in results:
        grouped[str(row.get(field, ""))].append(row)

    summaries = []

    for group_name, rows in sorted(grouped.items()):
        count = len(rows)

        exact = sum(row["title_exact_match"] for row in rows)
        practical_title = sum(row["title_practical_match"] for row in rows)

        reliable_loc_rows = [
            row for row in rows if row["location_reference_reliable"]
        ]

        practical_location = sum(
            row["location_practical_match"] for row in reliable_loc_rows
        )

        summaries.append({
            "group": group_name,
            "count": count,
            "title_exact_rate_percent": (
                exact / count * 100 if count else 0
            ),
            "title_practical_rate_percent": (
                practical_title / count * 100 if count else 0
            ),
            "location_reliable_n": len(reliable_loc_rows),
            "location_practical_rate_percent": (
                practical_location / len(reliable_loc_rows) * 100
                if reliable_loc_rows else None
            ),
            "avg_title_character_similarity": average([
                parse_decimal(row["title_character_similarity"])
                for row in rows
            ]),
            "avg_title_word_similarity": average([
                parse_decimal(row["title_word_similarity_f1"])
                for row in rows
            ]),
        })

    return summaries


# ============================================================
# Evaluation
# ============================================================

def evaluate(
    groundtruth_path,
    predictions_path,
    output_path,
    summary_path,
    files_dir=None,
    title_char_threshold=0.92,
    title_word_threshold=0.90,
    location_char_threshold=0.85,
    location_word_threshold=0.80,
):
    gt_rows = read_groundtruth(groundtruth_path)
    pred_rows = read_csv(predictions_path)

    require_columns(
        gt_rows,
        ["filename", "actual_title", "actual_location", "location_status"],
        "Ground-truth file",
    )

    require_columns(
        pred_rows,
        ["Filename", "Title"],
        "Prediction CSV",
    )

    pred_lookup = {}

    for row in pred_rows:
        key = normalize_filename(row.get("Filename"))

        if not key:
            continue

        if key in pred_lookup:
            raise ValueError(
                f"Duplicate prediction filename: {row.get('Filename')}"
            )

        pred_lookup[key] = row

    file_lookup = build_file_lookup(files_dir)

    results = []
    matched_prediction_keys = set()

    for gt in gt_rows:
        filename = gt.get("filename", "")
        actual_title = gt.get("actual_title", "") or ""
        actual_location = gt.get("actual_location", "") or ""
        location_status = gt.get("location_status", "") or ""

        key = normalize_filename(filename)
        pred = pred_lookup.get(key)

        if pred is None:
            pred = {}
            prediction_found = False
            predicted_title = ""
            predicted_location = ""
        else:
            prediction_found = True
            matched_prediction_keys.add(key)
            predicted_title = pred.get("Title", "") or ""
            predicted_location = pred.get("Location", "") or ""

        # ---------- title ----------
        title_exact = (
            prediction_found and predicted_title == actual_title
        )

        title_char = character_similarity(
            actual_title, predicted_title
        )

        (
            title_word,
            title_missing_words,
            title_extra_words,
        ) = word_metrics(actual_title, predicted_title)

        title_practical = practical_title_match(
            title_exact,
            title_char,
            title_word,
            title_char_threshold,
            title_word_threshold,
        )

        title_category = suggested_title_error_category(
            actual_title,
            predicted_title,
            title_exact,
            title_char,
            title_word,
            title_missing_words,
            title_extra_words,
        )

        # ---------- location ----------
        location_exact = (
            prediction_found
            and predicted_location == actual_location
            and actual_location != ""
        )

        location_char = character_similarity(
            actual_location, predicted_location
        )

        (
            location_word,
            location_missing_words,
            location_extra_words,
        ) = word_metrics(actual_location, predicted_location)

        location_practical = practical_location_match(
            actual_location,
            predicted_location,
            location_exact,
            location_char,
            location_word,
            location_char_threshold,
            location_word_threshold,
        )

        loc_quality = location_reference_quality(location_status)
        loc_reliable = location_reference_reliable(location_status)

        # ---------- combined retrieval ----------
        if not prediction_found:
            retrieval_status = "MISSING_PREDICTION"
        elif loc_quality == "UNCERTAIN":
            # Do not let an uncertain human location decide pass/fail.
            retrieval_status = (
                "TITLE_USABLE_LOCATION_REFERENCE_UNCERTAIN"
                if title_practical
                else "TITLE_NOT_USABLE_LOCATION_REFERENCE_UNCERTAIN"
            )
        elif title_practical and location_practical:
            retrieval_status = "FULLY_IDENTIFIED"
        elif title_practical and not location_practical:
            retrieval_status = "TITLE_ONLY"
        elif not title_practical and location_practical:
            retrieval_status = "LOCATION_ONLY"
        else:
            retrieval_status = "FAILED"

        # ---------- process metadata ----------
        title_probability = parse_decimal(
            pred.get("Title_probability")
        )
        scale_percent = parse_decimal(
            pred.get("Scale_percent")
        )
        prediction_original_mp = parse_decimal(
            pred.get("Original_size_MP")
        )

        # ---------- intrinsic metadata ----------
        file_meta = get_file_metadata(filename, file_lookup)
        intrinsic_mp = parse_decimal(
            file_meta.get("image_megapixels")
        )
        mp_for_band = (
            intrinsic_mp
            if intrinsic_mp is not None
            else prediction_original_mp
        )
        dpi_x = parse_decimal(file_meta.get("dpi_x"))

        features = title_features(actual_title)

        result = {
            "filename": filename,
            "prediction_found": prediction_found,

            # Ground truth
            "actual_title": actual_title,
            "actual_location": actual_location,
            "location_status": location_status,
            "location_reference_quality": loc_quality,
            "location_reference_reliable": loc_reliable,

            # Predictions
            "predicted_title": predicted_title,
            "predicted_location": predicted_location,

            # Strict title metrics
            "title_exact_match": title_exact,
            "title_character_similarity": round(title_char, 4),
            "title_word_similarity_f1": round(title_word, 4),

            # Practical title metric
            "title_practical_match": title_practical,
            "title_practical_threshold_char": title_char_threshold,
            "title_practical_threshold_word": title_word_threshold,

            # Title detail
            "title_missing_words": " | ".join(title_missing_words),
            "title_extra_words": " | ".join(title_extra_words),
            "title_suggested_error_category": title_category,
            "actual_title_length": len(actual_title),
            "predicted_title_length": len(predicted_title),
            "actual_title_word_count": len(tokenize(actual_title)),
            "predicted_title_word_count": len(tokenize(predicted_title)),
            **features,

            # Strict location metrics
            "location_exact_match": location_exact,
            "location_character_similarity": round(location_char, 4),
            "location_word_similarity_f1": round(location_word, 4),

            # Practical location metric
            "location_practical_match": location_practical,
            "location_practical_threshold_char": location_char_threshold,
            "location_practical_threshold_word": location_word_threshold,
            "location_core_containment": location_core_containment(
                actual_location, predicted_location
            ),
            "location_missing_words": " | ".join(location_missing_words),
            "location_extra_words": " | ".join(location_extra_words),

            # Combined utility
            "retrieval_status": retrieval_status,

            # Prediction/process metadata
            "title_confidence": pred.get("Title_confidence", "") or "",
            "title_probability": safe_round(title_probability, 4),
            "prediction_original_size_px": pred.get("Original_size_px", "") or "",
            "prediction_original_size_mp": pred.get("Original_size_MP", "") or "",
            "prediction_crop_size_px": pred.get("Crop_size_px", "") or "",
            "prediction_sent_size_px": pred.get("Sent_size_px", "") or "",
            "prediction_scale_factor": pred.get("Scale_factor", "") or "",
            "scale_percent": safe_round(scale_percent, 2),
            "prediction_sent_kb": pred.get("Sent_KB", "") or "",
            "preparation_time_s": pred.get("Preparation_time_s", "") or "",
            "first_response_time_s": pred.get("First_response_time_s", "") or "",
            "model_time_s": pred.get("Model_time_s", "") or "",
            "total_time_s": pred.get("Total_time_s", "") or "",

            # Grouping bands
            "title_length_band": title_length_band(len(actual_title)),
            "scale_band": scale_band(scale_percent),
            "dpi_band": dpi_band(dpi_x),
            "megapixel_band": megapixel_band(mp_for_band),
        }

        result.update(file_meta)
        results.append(result)

    extra_prediction_keys = set(pred_lookup) - matched_prediction_keys

    # --------------------------------------------------------
    # Detailed CSV
    # --------------------------------------------------------
    fieldnames = list(results[0].keys()) if results else []

    with open(
        output_path,
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as f:
        writer = csv.DictWriter(
            f,
            fieldnames=fieldnames,
            delimiter=";",
        )
        writer.writeheader()
        writer.writerows(results)

    # --------------------------------------------------------
    # Aggregate summary
    # --------------------------------------------------------
    total = len(results)

    exact_title_n = sum(r["title_exact_match"] for r in results)
    practical_title_n = sum(r["title_practical_match"] for r in results)

    reliable_location_rows = [
        r for r in results if r["location_reference_reliable"]
    ]
    strong_location_rows = [
        r for r in results
        if r["location_reference_quality"] == "STRONG"
    ]
    derived_location_rows = [
        r for r in results
        if r["location_reference_quality"] == "DERIVED"
    ]
    uncertain_location_rows = [
        r for r in results
        if r["location_reference_quality"] == "UNCERTAIN"
    ]

    exact_location_n = sum(
        r["location_exact_match"] for r in reliable_location_rows
    )
    practical_location_n = sum(
        r["location_practical_match"] for r in reliable_location_rows
    )

    retrieval_counts = defaultdict(int)
    for r in results:
        retrieval_counts[r["retrieval_status"]] += 1

    error_counts = defaultdict(int)
    for r in results:
        error_counts[r["title_suggested_error_category"]] += 1

    summary_rows = [
        [
            "section", "metric", "value", "count",
            "title_exact_rate_percent",
            "title_practical_rate_percent",
            "location_practical_rate_percent",
            "avg_title_character_similarity",
            "avg_title_word_similarity"
        ],
        ["overall", "files", total, "", "", "", "", "", ""],
        [
            "overall",
            "title_exact_matches",
            exact_title_n,
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "overall",
            "title_exact_rate_percent",
            round(exact_title_n / total * 100, 2) if total else 0,
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "overall",
            "title_practical_matches",
            practical_title_n,
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "overall",
            "title_practical_rate_percent",
            round(practical_title_n / total * 100, 2) if total else 0,
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "overall",
            "avg_title_character_similarity",
            safe_round(average([
                r["title_character_similarity"] for r in results
            ]), 4),
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "overall",
            "avg_title_word_similarity",
            safe_round(average([
                r["title_word_similarity_f1"] for r in results
            ]), 4),
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "location",
            "reliable_location_references",
            len(reliable_location_rows),
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "location",
            "strong_manual_references",
            len(strong_location_rows),
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "location",
            "derived_references",
            len(derived_location_rows),
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "location",
            "uncertain_references_excluded_from_main_rate",
            len(uncertain_location_rows),
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "location",
            "location_exact_rate_reliable_percent",
            round(
                exact_location_n / len(reliable_location_rows) * 100, 2
            ) if reliable_location_rows else "",
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "location",
            "location_practical_rate_reliable_percent",
            round(
                practical_location_n / len(reliable_location_rows) * 100, 2
            ) if reliable_location_rows else "",
            "",
            "",
            "",
            "",
            "",
            "",
        ],
    ]

    for category, count in sorted(error_counts.items()):
        summary_rows.append([
            "title_error_category", category, count,
            "", "", "", "", "", ""
        ])

    for status, count in sorted(retrieval_counts.items()):
        summary_rows.append([
            "retrieval_status", status, count,
            "", "", "", "", "", ""
        ])

    grouping_fields = [
        ("file_type", "file_extension"),
        ("title_length_band", "title_length_band"),
        ("scale_band", "scale_band"),
        ("dpi_band", "dpi_band"),
        ("megapixel_band", "megapixel_band"),
        ("location_reference_quality", "location_reference_quality"),
    ]

    for section, field in grouping_fields:
        for group in summarize_groups(results, field):
            summary_rows.append([
                section,
                group["group"],
                "",
                group["count"],
                safe_round(group["title_exact_rate_percent"], 2),
                safe_round(group["title_practical_rate_percent"], 2),
                safe_round(group["location_practical_rate_percent"], 2),
                safe_round(group["avg_title_character_similarity"], 4),
                safe_round(group["avg_title_word_similarity"], 4),
            ])

    with open(
        summary_path,
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerows(summary_rows)

    # --------------------------------------------------------
    # Console report
    # --------------------------------------------------------
    print()
    print("================================================")
    print("TITLE + LOCATION EVALUATION - V4")
    print("================================================")
    print(f"Files evaluated:                 {total}")
    print(f"Prediction rows:                {len(pred_rows)}")
    print(f"Missing predictions:            {sum(not r['prediction_found'] for r in results)}")
    print(f"Extra predictions:              {len(extra_prediction_keys)}")

    print()
    print("TITLE")
    print("------------------------------------------------")
    print(
        f"Exact title matches:            "
        f"{exact_title_n}/{total} "
        f"({exact_title_n / total * 100:.2f}%)"
        if total else "Exact title matches:            0/0"
    )
    print(
        f"Practical title matches:        "
        f"{practical_title_n}/{total} "
        f"({practical_title_n / total * 100:.2f}%)"
        if total else "Practical title matches:        0/0"
    )
    print(
        f"Average character similarity:   "
        f"{average([r['title_character_similarity'] for r in results]):.4f}"
        if results else ""
    )
    print(
        f"Average word similarity:        "
        f"{average([r['title_word_similarity_f1'] for r in results]):.4f}"
        if results else ""
    )
    print(
        f"Practical thresholds:           "
        f"char >= {title_char_threshold:.2f} OR "
        f"word >= {title_word_threshold:.2f}"
    )

    print()
    print("LOCATION")
    print("------------------------------------------------")
    print(
        f"Reliable location references:   "
        f"{len(reliable_location_rows)}/{total}"
    )
    print(
        f"  manually identified:          {len(strong_location_rows)}"
    )
    print(
        f"  derived from title:           {len(derived_location_rows)}"
    )
    print(
        f"Uncertain human references:     {len(uncertain_location_rows)} "
        "(excluded from main rate)"
    )

    if reliable_location_rows:
        print(
            f"Exact location matches:         "
            f"{exact_location_n}/{len(reliable_location_rows)} "
            f"({exact_location_n / len(reliable_location_rows) * 100:.2f}%)"
        )
        print(
            f"Practical location matches:     "
            f"{practical_location_n}/{len(reliable_location_rows)} "
            f"({practical_location_n / len(reliable_location_rows) * 100:.2f}%)"
        )

    print(
        f"Practical thresholds:           "
        f"char >= {location_char_threshold:.2f} OR "
        f"word >= {location_word_threshold:.2f} "
        f"OR core-location containment"
    )

    print()
    print("COMBINED RETRIEVAL STATUS")
    print("------------------------------------------------")
    for status, count in sorted(
        retrieval_counts.items(),
        key=lambda item: (-item[1], item[0]),
    ):
        print(f"{status:43} {count}")

    print()
    print("TITLE ERROR CATEGORIES")
    print("------------------------------------------------")
    for category, count in sorted(
        error_counts.items(),
        key=lambda item: (-item[1], item[0]),
    ):
        print(f"{category:43} {count}")

    if files_dir:
        found = sum(r["source_file_found"] for r in results)
        print()
        print(f"Original source files found:    {found}/{total}")

    print()
    print(f"Detailed results: {output_path}")
    print(f"Summary:          {summary_path}")


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Evaluate title and location recognition for technical drawings."
        )
    )

    parser.add_argument(
        "groundtruth",
        help="Ground-truth .csv or .xlsx file",
    )

    parser.add_argument(
        "predictions",
        help="Prediction CSV",
    )

    parser.add_argument(
        "files_dir",
        nargs="?",
        default=None,
        help="Optional directory containing original TIF/TIFF/PDF drawings",
    )

    parser.add_argument(
        "-o",
        "--output",
        default="evaluation_results_v4.csv",
        help="Detailed output CSV",
    )

    parser.add_argument(
        "-s",
        "--summary",
        default="evaluation_summary_v4.csv",
        help="Summary output CSV",
    )

    parser.add_argument(
        "--title-char-threshold",
        type=float,
        default=0.92,
        help="Practical title character-similarity threshold (default 0.92)",
    )

    parser.add_argument(
        "--title-word-threshold",
        type=float,
        default=0.90,
        help="Practical title word-F1 threshold (default 0.90)",
    )

    parser.add_argument(
        "--location-char-threshold",
        type=float,
        default=0.85,
        help="Practical location character threshold (default 0.85)",
    )

    parser.add_argument(
        "--location-word-threshold",
        type=float,
        default=0.80,
        help="Practical location word-F1 threshold (default 0.80)",
    )

    args = parser.parse_args()

    evaluate(
        Path(args.groundtruth),
        Path(args.predictions),
        Path(args.output),
        Path(args.summary),
        args.files_dir,
        args.title_char_threshold,
        args.title_word_threshold,
        args.location_char_threshold,
        args.location_word_threshold,
    )


if __name__ == "__main__":
    main()
