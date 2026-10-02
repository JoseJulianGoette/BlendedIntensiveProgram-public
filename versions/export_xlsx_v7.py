"""
FROZEN Excel export of v7 (with date). Current export: export_xlsx.py (v8+).

Formatted Excel export of the analysis results.

Sheets:
    Plans      one row per plan, traffic lights as cell colours,
               review columns for the human check (human in the loop)
    Summary    counts per traffic light and field
    Run        settings of the run (model, prompt, thresholds ...)
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from models_v7 import ImageData, TrafficLight


INK = "1C2440"
BLUE = "34459A"
LINE = "C8CED8"
PAPER = "EEF1F4"

LEVEL_STYLE = {
    TrafficLight.GREEN: ("confident", "CDE9D6", "17613A"),
    TrafficLight.YELLOW: ("check", "FFE3B8", "7A4A00"),
    TrafficLight.RED: ("uncertain", "F4CFCB", "8E2119"),
    None: ("–", "FFFFFF", "6B7385"),
}

LEVEL_ORDER = [TrafficLight.RED, TrafficLight.YELLOW, TrafficLight.GREEN, None]

REVIEW_STATES = ["open", "confirmed", "corrected", "not a plan"]

FIELDS = [("location", "Location"), ("title", "Title"), ("date", "Date")]


# (header, width, is review column)
COLUMNS = (
    [("Overall", 11, False), ("File", 34, False), ("Note in file name", 16, False)]
    + [
        col
        for _, label in FIELDS
        for col in (
            (label, 48 if label == "Title" else 24, False),
            (f"{label} rating", 11, False),
            (f"{label} %", 7, False),
        )
    ]
    + [
        ("Error", 24, False),
        ("Review status", 13, True),
        ("Corrected location", 22, True),
        ("Corrected title", 34, True),
        ("Corrected date", 14, True),
        ("Comment", 30, True),
    ]
)


thin = Side(style="thin", color=LINE)
BORDER = Border(bottom=thin)


def level_cell(cell, level):
    text, fill, font = LEVEL_STYLE[level]
    cell.value = text
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(color=font, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="top")


def header_row(ws, headers, review_flags=None):
    for i, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=text)
        review = review_flags[i - 1] if review_flags else False
        cell.fill = PatternFill("solid", fgColor=BLUE if review else INK)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30


def write_plans_sheet(ws, entries: list[ImageData], folder: Path | None):

    header_row(ws, [c[0] for c in COLUMNS], [c[2] for c in COLUMNS])

    for i, (_, width, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Red first: what needs a human comes on top
    entries = sorted(
        entries,
        key=lambda e: (LEVEL_ORDER.index(e.overall_level), e.filename.casefold())
    )

    for r, e in enumerate(entries, start=2):

        level_cell(ws.cell(row=r, column=1), e.overall_level)

        name = ws.cell(row=r, column=2, value=e.filename)
        if folder is not None:
            name.hyperlink = (folder / e.filename).as_uri()
            name.font = Font(color=BLUE, underline="single")

        ws.cell(row=r, column=3, value=", ".join(e.filename_flags) or None)

        col = 4
        for field, _ in FIELDS:
            confidence = getattr(e, f"{field}_confidence")
            ws.cell(row=r, column=col, value=getattr(e.metadata, field))
            level_cell(ws.cell(row=r, column=col + 1), confidence.level if confidence else None)
            if confidence and confidence.probability is not None:
                pct = ws.cell(row=r, column=col + 2, value=confidence.probability)
                pct.number_format = "0%"
            col += 3

        ws.cell(row=r, column=col, value=e.error)
        ws.cell(row=r, column=col + 1, value="open")

        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if c != 1 and "rating" not in COLUMNS[c - 1][0]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    last_row = max(2, len(entries) + 1)
    last_col = get_column_letter(len(COLUMNS))

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"

    status_col = get_column_letter([c[0] for c in COLUMNS].index("Review status") + 1)
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(REVIEW_STATES) + '"',
        allow_blank=True
    )
    ws.add_data_validation(validation)
    validation.add(f"{status_col}2:{status_col}{last_row}")


def write_overview_sheet(ws, entries: list[ImageData]):

    ws.column_dimensions["A"].width = 18
    for c in "BCDE":
        ws.column_dimensions[c].width = 11

    header_row(ws, ["Field", "confident", "check", "uncertain", "no rating"])

    # Header cells in the colour of their traffic light
    for col, level in zip(range(2, 6), [TrafficLight.GREEN, TrafficLight.YELLOW, TrafficLight.RED, None]):
        _, fill, font = LEVEL_STYLE[level]
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=fill)
        ws.cell(row=1, column=col).font = Font(color=font, bold=True)

    rows = [
        (label, [getattr(e, f"{field}_confidence") for e in entries])
        for field, label in FIELDS
    ]

    for r, (label, confidences) in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        levels = [c.level if c else None for c in confidences]
        for col, level in zip(range(2, 6), [TrafficLight.GREEN, TrafficLight.YELLOW, TrafficLight.RED, None]):
            ws.cell(row=r, column=col, value=levels.count(level))

    r = len(rows) + 2
    ws.cell(row=r, column=1, value="Overall").font = Font(bold=True)
    overall = [e.overall_level for e in entries]
    for col, level in zip(range(2, 6), [TrafficLight.GREEN, TrafficLight.YELLOW, TrafficLight.RED, None]):
        ws.cell(row=r, column=col, value=overall.count(level)).font = Font(bold=True)

    info = [
        ("Plans", len(entries)),
        ("with errors", sum(1 for e in entries if e.error)),
        ("with note in file name", sum(1 for e in entries if e.filename_flags)),
    ]

    for i, (label, value) in enumerate(info, start=r + 2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)

    legend_row = r + 2 + len(info) + 1
    legend = [
        "Overall = worst rating of location and title (the search fields).",
        "confident = the model is sure, check = please review,",
        "uncertain = probably wrong or not found on the plan.",
        "All plans are included, also uncertain ones and errors.",
    ]
    for i, text in enumerate(legend):
        ws.cell(row=legend_row + i, column=1, value=text).font = Font(color="6B7385")


def write_run_sheet(ws, run_info: dict):

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100

    header_row(ws, ["Setting", "Value"])

    for r, (key, value) in enumerate(run_info.items(), start=2):
        ws.cell(row=r, column=1, value=key).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=str(value))
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_xlsx(
    entries: list[dict | ImageData],
    path: Path,
    run_info: dict | None = None,
    folder: Path | None = None
) -> Path:
    """Write the formatted workbook and return its path."""

    entries = [
        e if isinstance(e, ImageData) else ImageData.model_validate(e)
        for e in entries
    ]

    wb = Workbook()

    plans = wb.active
    plans.title = "Plans"
    write_plans_sheet(plans, entries, folder)

    write_overview_sheet(wb.create_sheet("Summary"), entries)

    if run_info:
        write_run_sheet(wb.create_sheet("Run"), run_info)

    wb.save(path)

    return path
