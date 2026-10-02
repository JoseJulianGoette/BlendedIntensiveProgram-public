"""
Formatted Excel export of the analysis results (v8+).

Sheets:
    Plans      one row per plan: final values (reviewed if corrected),
               traffic lights as cell colours, review status,
               and what the model originally read
    Summary    counts per traffic light, field and review status
    Run        settings of the run (model, prompt, thresholds ...)

v7 (with date): see export_xlsx_v7.py.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models import FIELDS, LEVEL_ORDER, ImageData, TrafficLight


INK = "1C2440"
GREY = "555555"
LINE = "C8CED8"

LEVEL_STYLE = {
    TrafficLight.GREEN: ("confident", "CDE9D6", "17613A"),
    TrafficLight.YELLOW: ("check", "FFE3B8", "7A4A00"),
    TrafficLight.RED: ("uncertain", "F4CFCB", "8E2119"),
    None: ("–", "FFFFFF", "6B7385"),
}

STATUS_STYLE = {
    "open": ("FFFFFF", "6B7385"),
    "confirmed": ("E2EAF6", "2B4DA3"),
    "corrected": ("E2EAF6", "2B4DA3"),
}

LABEL = {"location": "Location", "title": "Title"}


# (header, width, group) - group "model" = what the model read
COLUMNS = (
    [("Overall", 11, None), ("File", 34, None), ("Note in file name", 16, None)]
    + [
        col
        for field in FIELDS
        for col in (
            (LABEL[field], 48 if field == "title" else 26, None),
            (f"{LABEL[field]} rating", 11, None),
        )
    ]
    + [("Review status", 13, None), ("Reviewed at", 17, None)]
    + [
        col
        for field in FIELDS
        for col in (
            (f"Model {LABEL[field].lower()}", 40 if field == "title" else 24, "model"),
            (f"Model {LABEL[field].lower()} %", 9, "model"),
        )
    ]
    + [("Error", 30, None)]
)


BORDER = Border(bottom=Side(style="thin", color=LINE))


def level_cell(cell, level):
    text, fill, font = LEVEL_STYLE[level]
    cell.value = text
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(color=font, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="top")


def header_row(ws, headers, groups=None):
    for i, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=text)
        model = groups is not None and groups[i - 1] == "model"
        cell.fill = PatternFill("solid", fgColor=GREY if model else INK)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30


def write_plans_sheet(ws, entries: list[ImageData], folder):

    header_row(ws, [c[0] for c in COLUMNS], [c[2] for c in COLUMNS])

    for i, (_, width, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Uncertain first: what needs a human comes on top
    rank = LEVEL_ORDER + [None]
    entries = sorted(entries, key=lambda e: (rank.index(e.overall_level), e.filename.casefold()))

    for r, e in enumerate(entries, start=2):

        values = [None, e.filename, ", ".join(e.filename_flags) or None]
        for field in FIELDS:
            values += [e.value(field), None]
        values += [e.review_status, e.review.reviewed_at if e.review else None]
        for field in FIELDS:
            confidence = getattr(e, f"{field}_confidence")
            values += [
                getattr(e.metadata, field),
                confidence.probability if confidence else None,
            ]
        values.append(e.error)

        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if COLUMNS[c - 1][0].endswith("%") and value is not None:
                cell.number_format = "0%"

        level_cell(ws.cell(row=r, column=1), e.overall_level)

        for i, field in enumerate(FIELDS):
            level_cell(ws.cell(row=r, column=5 + 2 * i), e.level(field))

        status = ws.cell(row=r, column=4 + 2 * len(FIELDS))
        fill, font = STATUS_STYLE[e.review_status]
        status.fill = PatternFill("solid", fgColor=fill)
        status.font = Font(color=font, bold=True)

        if folder is not None:
            name = ws.cell(row=r, column=2)
            name.hyperlink = (folder / e.filename).as_uri()
            name.font = Font(color="2B4DA3", underline="single")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(2, len(entries) + 1)}"


def write_summary_sheet(ws, entries: list[ImageData]):

    ws.column_dimensions["A"].width = 24
    for c in "BCDE":
        ws.column_dimensions[c].width = 12

    levels = [TrafficLight.GREEN, TrafficLight.YELLOW, TrafficLight.RED, None]

    header_row(ws, ["Field"] + [LEVEL_STYLE[level][0] for level in levels[:3]] + ["no rating"])

    for col, level in enumerate(levels, start=2):
        _, fill, font = LEVEL_STYLE[level]
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=fill)
        ws.cell(row=1, column=col).font = Font(color=font, bold=True)

    rows = [(LABEL[f], [e.level(f) for e in entries]) for f in FIELDS]
    rows.append(("Overall", [e.overall_level for e in entries]))

    for r, (label, values) in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        for col, level in enumerate(levels, start=2):
            ws.cell(row=r, column=col, value=values.count(level))

    info = [
        ("Plans", len(entries)),
        ("reviewed", sum(1 for e in entries if e.review)),
        ("  of which corrected", sum(1 for e in entries if e.review_status == "corrected")),
        ("not reviewed yet", sum(1 for e in entries if not e.review)),
        ("with errors", sum(1 for e in entries if e.error)),
        ("with note in file name", sum(1 for e in entries if e.filename_flags)),
    ]

    start = len(rows) + 3
    for i, (label, value) in enumerate(info):
        ws.cell(row=start + i, column=1, value=label)
        ws.cell(row=start + i, column=2, value=value)

    legend = [
        "Overall = worst rating of location and title.",
        "Ratings and values include human corrections from the review.",
        "confident = sure, check = please review,",
        "uncertain = probably wrong or not found on the plan.",
        "All plans are included, also uncertain ones and errors.",
    ]
    for i, text in enumerate(legend):
        ws.cell(row=start + len(info) + 1 + i, column=1, value=text).font = Font(color="6B7385")


def write_run_sheet(ws, run_info: dict):

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100

    header_row(ws, ["Setting", "Value"])

    for r, (key, value) in enumerate(run_info.items(), start=2):
        ws.cell(row=r, column=1, value=key).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=str(value))
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_xlsx(entries: list, path, run_info: dict | None = None, folder=None):
    """Write the formatted workbook. `path` may be a file path or a file object."""

    entries = [
        e if isinstance(e, ImageData) else ImageData.model_validate(e)
        for e in entries
    ]

    wb = Workbook()

    plans = wb.active
    plans.title = "Plans"
    write_plans_sheet(plans, entries, folder)

    write_summary_sheet(wb.create_sheet("Summary"), entries)

    if run_info:
        write_run_sheet(wb.create_sheet("Run"), run_info)

    wb.save(path)

    return path
