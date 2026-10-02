from pathlib import Path
from openpyxl import load_workbook
import matplotlib.pyplot as plt


# ============================================================
# PATHS
# ============================================================

# This script is located at:
# repo/analysis/plot_ground_truth_location_status.py
ANALYSIS_DIR = Path(__file__).resolve().parent

# Input Excel file:
# repo/analysis/results/title_location_groundtruth_vs_predictions_v9.xlsx
INPUT_FILE = (
    ANALYSIS_DIR
    / "results"
    / "title_location_groundtruth_vs_predictions_v9.xlsx"
)

# Central chart folder:
# repo/analysis/results/charts/
OUTPUT_DIR = (
    ANALYSIS_DIR
    / "results"
    / "charts"
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)

# Output chart:
OUTPUT_CHART = (
    OUTPUT_DIR
    / "ground_truth_location_status.png"
)

# ============================================================
# BVG-INSPIRED COLORS
# ============================================================

BVG_YELLOW = "#F7B500"
BVG_BLACK = "#111111"
BVG_GRAY = "#A7A7A7"


# ============================================================
# CHECK INPUT FILE
# ============================================================

print()
print("========================================")
print("GROUND-TRUTH LOCATION VISUALIZATION")
print("========================================")

print()
print("Reading ground-truth data from:")
print(INPUT_FILE)

print()
print("Chart will be saved to:")
print(OUTPUT_DIR)


if not INPUT_FILE.exists():
    raise FileNotFoundError(
        "\nCould not find the Excel file at:\n"
        f"{INPUT_FILE}"
    )


# ============================================================
# READ EXCEL
# ============================================================

workbook = load_workbook(
    INPUT_FILE,
    data_only=True
)

sheet = workbook["Comparison"]


# Read column names
headers = [
    cell.value
    for cell in sheet[1]
]


# ============================================================
# COUNTERS
# ============================================================

manually_identified = 0
derived_from_title = 0
uncertain_by_human = 0


# ============================================================
# READ ROWS
# ============================================================

for values in sheet.iter_rows(
    min_row=2,
    values_only=True
):

    row = dict(
        zip(headers, values)
    )

    status = row.get(
        "location_status"
    )

    if status is None:
        continue

    status = (
        str(status)
        .strip()
        .lower()
    )


    if status == "manually_identified":

        manually_identified += 1


    elif status == "derived_from_title":

        derived_from_title += 1


    elif status == "uncertain_by_human":

        uncertain_by_human += 1


# ============================================================
# SUMMARY VALUES
# ============================================================

total = (
    manually_identified
    + derived_from_title
    + uncertain_by_human
)

known_location = (
    manually_identified
    + derived_from_title
)


# ============================================================
# CREATE CHART
# ============================================================

categories = [
    "Manually\nidentified",
    "Derived from\nreference title",
    "Uncertain\nby human"
]

counts = [
    manually_identified,
    derived_from_title,
    uncertain_by_human
]

colors = [
    BVG_YELLOW,
    BVG_BLACK,
    BVG_GRAY
]


plt.figure(
    figsize=(9, 6)
)

bars = plt.bar(
    categories,
    counts,
    color=colors,
    edgecolor=BVG_BLACK,
    linewidth=1.2,
    width=0.6
)


# ============================================================
# ADD VALUES ABOVE BARS
# ============================================================

for bar in bars:

    height = bar.get_height()

    plt.text(
        bar.get_x()
        + bar.get_width() / 2,
        height + 0.3,
        str(int(height)),
        ha="center",
        va="bottom",
        fontsize=12,
        fontweight="bold"
    )


# ============================================================
# LABELS
# ============================================================

plt.ylabel(
    "Number of plans",
    fontsize=11
)

plt.title(
    "Ground-Truth Location Availability",
    fontsize=15,
    fontweight="bold",
    pad=15
)


# ============================================================
# STYLE
# ============================================================

plt.grid(
    axis="y",
    linestyle="--",
    alpha=0.25
)

plt.gca().spines["top"].set_visible(False)
plt.gca().spines["right"].set_visible(False)

plt.tight_layout()


# ============================================================
# SAVE CHART
# ============================================================

plt.savefig(
    OUTPUT_CHART,
    dpi=300,
    bbox_inches="tight"
)

plt.close()


# ============================================================
# SUMMARY
# ============================================================

print()
print("----------------------------------------")
print("GROUND-TRUTH LOCATION SUMMARY")
print("----------------------------------------")

print(f"Total plans:                    {total}")
print(f"Reference location available:   {known_location}")
print(f"Uncertain by human:             {uncertain_by_human}")

print()
print("Location annotation source:")
print(f"  Manually identified:          {manually_identified}")
print(f"  Derived from reference title: {derived_from_title}")
print(f"  Uncertain by human:           {uncertain_by_human}")

print()
print("Chart created:")
print(OUTPUT_CHART)